import os
import sys

import pytest

from semsearch.config import Config
from semsearch.extract.office import DocxExtractor, PptxExtractor, XlsxExtractor
from semsearch.extract.registry import ExtractorRegistry, build_default_registry
from semsearch.extract.text import TextExtractor, decode_bytes
from semsearch.models import ExtractResult

WIN = sys.platform == "win32"


def test_decode_bom_and_fallbacks():
    assert decode_bytes("héllo".encode("utf-8"))[1] == "utf-8"
    assert decode_bytes(b"\xef\xbb\xbfabc")[0] == "abc"
    t, enc = decode_bytes("hi there".encode("utf-16"))
    assert t == "hi there" and enc == "utf-16"
    t, enc = decode_bytes(b"caf\xe9")
    assert t == "café" and enc == "cp1252"
    with pytest.raises(ValueError):
        decode_bytes(b"\x00\x01\x02\x03" * 10)


def test_text_extractor_statuses(tmp_path):
    ex = TextExtractor({".txt"}, max_chars=10)
    p = tmp_path / "a.txt"
    p.write_text("hello world this is long", encoding="utf-8")
    r = ex.extract(str(p), ".txt")
    assert r.status == "ok" and r.text == "hello worl"
    (tmp_path / "e.txt").write_bytes(b"")
    assert ex.extract(str(tmp_path / "e.txt"), ".txt").status == "empty"
    (tmp_path / "b.txt").write_bytes(b"\x00" * 50)
    assert ex.extract(str(tmp_path / "b.txt"), ".txt").status == "binary"


class Failing:
    name = "failing"

    def supports(self, ext):
        return True

    def extract(self, path, ext):
        raise RuntimeError("kaboom")


class Empty:
    name = "empty"

    def supports(self, ext):
        return True

    def extract(self, path, ext):
        return ExtractResult("", "empty", self.name)


class Good:
    name = "good"

    def supports(self, ext):
        return True

    def extract(self, path, ext):
        return ExtractResult("text!", "ok", self.name)


def test_registry_falls_through_and_records_attempts(tmp_path):
    p = tmp_path / "x.foo"
    p.write_text("x")
    reg = ExtractorRegistry({".foo": [Failing(), Empty(), Good()]})
    r = reg.extract(str(p))
    assert r.ok and r.method == "good"
    assert [a["extractor"] for a in r.meta["attempts"]] == ["failing", "empty", "good"]
    assert r.meta["attempts"][0]["status"] == "error" and "kaboom" in r.meta["attempts"][0]["error"]
    reg2 = ExtractorRegistry({".foo": [Failing(), Empty()]})
    r2 = reg2.extract(str(p))
    assert not r2.ok and r2.status == "empty"
    assert reg.extract(str(p), ".nope").status == "unsupported"
    assert reg.extract(str(tmp_path / "missing.foo")).status == "missing"


def test_docx_pptx_xlsx_fallback_extractors(tmp_path):
    import docx
    import openpyxl
    from pptx import Presentation

    d = docx.Document()
    d.add_paragraph("The blackboard design tracks reads.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "key", "value"
    d.save(tmp_path / "a.docx")
    r = DocxExtractor().extract(str(tmp_path / "a.docx"), ".docx")
    assert "blackboard design" in r.text and "key | value" in r.text

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Roadmap"
    s.placeholders[1].text = "Ship semantic search"
    prs.save(tmp_path / "a.pptx")
    r = PptxExtractor().extract(str(tmp_path / "a.pptx"), ".pptx")
    assert "Roadmap" in r.text and "Ship semantic search" in r.text and "# Slide 1" in r.text

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["item", "cost"])
    ws.append(["gpu", 1999])
    wb.save(tmp_path / "a.xlsx")
    r = XlsxExtractor().extract(str(tmp_path / "a.xlsx"), ".xlsx")
    assert "# Sheet: Budget" in r.text and "gpu | 1999" in r.text


def test_default_registry_covers_configured_types():
    cfg = Config()
    reg = build_default_registry(cfg)
    exts = reg.supported_extensions()
    for e in (".md", ".py", ".pdf", ".docx", ".pptx", ".xlsx", ".json", ".yaml"):
        assert e in exts
    assert reg.extractors_for(".pdf")[0].name == "pypdf"


@pytest.mark.skipif(not WIN, reason="IFilter is Windows-only")
def test_ifilter_extracts_docx(tmp_path):
    import docx
    from semsearch.extract.ifilter import IFilterExtractor, has_registered_filter
    if not has_registered_filter(".docx"):
        pytest.skip("no docx IFilter registered on this machine")
    d = docx.Document()
    d.add_paragraph("Receipts make irreversible operations auditable.")
    d.save(tmp_path / "f.docx")
    r = IFilterExtractor().extract(str(tmp_path / "f.docx"), ".docx")
    assert r.ok and "auditable" in r.text and r.method.startswith("ifilter")


@pytest.mark.skipif(not WIN, reason="Windows-only")
def test_isolated_extractor_survives_hang_and_restarts(tmp_path):
    from semsearch.extract.isolated import IsolatedExtractor
    cfg = Config(roots=[str(tmp_path)], data_dir=str(tmp_path / "d"))
    cfg.indexing.extract_timeout_s = 3.0
    reg = build_default_registry(cfg)
    iso = IsolatedExtractor(cfg, reg, timeout_s=3.0)
    try:
        import docx
        d = docx.Document()
        d.add_paragraph("hello from the child process")
        d.save(tmp_path / "c.docx")
        r = iso.extract(str(tmp_path / "c.docx"), ".docx")
        assert r.ok and "child process" in r.text
        # text files stay in-process
        (tmp_path / "t.txt").write_text("in process")
        assert iso.extract(str(tmp_path / "t.txt"), ".txt").text == "in process"
        # simulate a hang: kill the child and make sure the next call restarts it
        iso._proc.terminate()
        iso._proc.join(5)
        r2 = iso.extract(str(tmp_path / "c.docx"), ".docx")
        assert r2.ok and iso.restarts >= 1
    finally:
        iso.close()
